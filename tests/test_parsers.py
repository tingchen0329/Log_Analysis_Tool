"""
解析器回歸測試（只使用 tests/fixtures 的合成 Log，不含任何客戶資料）。

執行：python -m unittest discover -s tests -v
"""
import importlib.util
import os
import tempfile
import unittest

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIXTURES = os.path.join(ROOT, "tests", "fixtures")

_spec = importlib.util.spec_from_file_location("cla", os.path.join(ROOT, "Cisco_Log_Analysis.py"))
cla = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(cla)


def read_fixture(name):
    with open(os.path.join(FIXTURES, name), encoding="utf-8") as f:
        return f.read()


def analyze(*names):
    """只分析指定的 fixture 檔（放到暫存資料夾，模擬使用者的 Log 資料夾）。"""
    with tempfile.TemporaryDirectory() as d:
        for n in names:
            with open(os.path.join(FIXTURES, n), encoding="utf-8") as src, \
                 open(os.path.join(d, n), "w", encoding="utf-8") as dst:
                dst.write(src.read())
        devices, aps, results = cla.analyze_folder(d, log=lambda _m: None)
    return devices, aps, results


class TestCommandIndex(unittest.TestCase):
    def test_cisco_abbreviation_and_errors(self):
        log = ("SW#sh ver\nCisco IOS Software, Version 17.9\n"
               "SW#show in\n% Ambiguous command: \"show in\"\n"
               "SW#show inventory\nNAME: \"1\"\n"
               "SW#show processes cpu | i CPU\nCPU utilization for five seconds: 1%/0%; five minutes: 3%\n"
               "SW#end\n")
        idx = cla.build_prompt_index(log)
        self.assertTrue(idx.get("show version").startswith("Cisco IOS"))
        self.assertEqual(idx.get("show inventory"), 'NAME: "1"')          # 跳過錯誤輸出
        self.assertIn("five minutes: 3%", idx.get("show processes cpu"))  # 管線過濾版本備援
        self.assertEqual(idx.get("show interfaces"), "")

    def test_pasted_command_list_is_ignored(self):
        log = "N9K# terminal length 0\n!\nshow version\nshow inventory\nN9K# show version\nNXOS: version 10.3(6)\nN9K# exit\n"
        self.assertIn("NXOS", cla.build_prompt_index(log).get("show version"))


class TestC9800HA(unittest.TestCase):
    def test_ha_pair_rows(self):
        devices, aps, _ = analyze("10.0.0.10@WLC-HA.log")
        self.assertEqual(len(devices), 2)
        active, standby = devices
        self.assertTrue(active["HA Role"].startswith("Active (Chassis 2)"))
        self.assertEqual(active["Serial Number"], "FCLTEST0002")
        self.assertEqual(active["OOB IP"], "10.0.0.10")                    # WMI primary，不是 RMI secondary
        self.assertIn("15838712K", active["Memory Usage"])                 # platform sorted 優先
        self.assertEqual(standby["Hostname"], "○ (HA)")
        self.assertEqual(standby["Serial Number"], "FCLTEST0001")
        self.assertEqual(len(aps), 1)
        self.assertEqual(aps[0]["source"], "10.0.0.10@WLC-HA.log")


class TestFortiGate(unittest.TestCase):
    def test_ha_cluster_merged_by_serial(self):
        devices, _, results = analyze("10.1.1.1@FGT-HQ-primary.log", "10.1.1.2@FGT-HQ-secondary.log")
        self.assertTrue(all(r.dev_type == "FortiGate" for r in results))
        self.assertEqual(len(devices), 2)                                  # 兩份 Log、兩台、不重複
        primary, secondary = devices
        self.assertEqual((primary["HA Role"], primary["Serial Number"]), ("Primary", "FG100FTK19000001"))
        self.assertEqual((secondary["HA Role"], secondary["Serial Number"]), ("Secondary", "FG100FTK19000002"))
        self.assertEqual(primary["Model"], "FortiGate-100F")
        self.assertEqual(primary["Firmware Version"], "v7.2.8 build1639")
        self.assertEqual(primary["OOB IP"], "10.1.1.1")
        self.assertEqual(primary["CPU Usage"], "4%")                       # 100 − 96% idle
        self.assertEqual(primary["Power Status"], "異常")                  # PS2 VIN alarm=1
        self.assertEqual(secondary["Power Status"], "正常")                # 取自 Secondary 自己的 Log
        self.assertEqual(primary["HA Status"], "OK｜in-sync")
        self.assertEqual(secondary["CPU Usage"], "1%")
        log = primary["System Log"].splitlines()
        self.assertEqual(log, ["2024-05-01 14:59:00 [warning] Power supply failure: Power supply 2 failure"])

    def test_backup_conf_only(self):
        devices, _, _ = analyze("FGT-BRANCH.conf")
        r = devices[0]
        self.assertEqual((r["Hostname"], r["Model"], r["Firmware Version"]), ("FGT-BRANCH", "FG60F", "v7.0.15 build0632"))
        self.assertEqual(r["OOB IP"], "192.168.10.1")
        self.assertEqual((r["CPU Usage"], r["Fan Status"], r["Power Status"]), ("N/A", "N/A", "N/A"))


class TestAruba(unittest.TestCase):
    def test_vsf_members_and_health(self):
        devices, _, _ = analyze("aruba_vsf_show_tech.txt")
        self.assertEqual([r["VSF Role"] for r in devices], ["Conductor", "Standby", "Not Present"])
        conductor, standby, _missing = devices
        self.assertEqual(conductor["Serial Number"], "SG00TEST01")
        self.assertEqual(standby["Serial Number"], "SG00TEST02")
        self.assertEqual(conductor["OOB IP"], "10.9.9.9")
        self.assertEqual(conductor["Fan Status"], "異常")
        self.assertEqual(conductor["Ports Up/Total"], "2 / 3")
        self.assertEqual(conductor["Core Dumps"], "1")
        self.assertEqual((conductor["Primary Image"], conductor["Secondary Image"]), ("FL.10.16.1020", "FL.10.13.1040"))
        self.assertIn("2/1-CPU", conductor["Max Temperature"])
        self.assertTrue(conductor["System Log"].startswith("2026-04-17T06:26:50"))   # 同訊息保留最後一次


class TestIOSStack(unittest.TestCase):
    """C9300 三台 Stack：show version 成員表格 + Switch 0N 分段（格式依 Cisco 官方 C9200/C9300 重開機排查文件）。"""

    def setUp(self):
        self.devices, _, self.results = analyze("192.0.2.10@CORE-STK.log")

    def test_detect_and_members(self):
        self.assertEqual(self.results[0].dev_type, "Switch")
        self.assertIn("Stack | 3", self.results[0].summary)
        self.assertEqual([r["Serial Number"] for r in self.devices], ["FOC0000TST1", "FOC0000TST2", "FOC0000TST3"])
        self.assertEqual([r["Model"] for r in self.devices], ["C9300-48P", "C9300-48P", "C9300-24P"])
        self.assertTrue(all(r["Firmware Version"] == "17.09.05" for r in self.devices))

    def test_member_mac_matched_by_serial(self):
        self.assertEqual([r["Ethernet MAC"] for r in self.devices],
                         ["aa:bb:cc:00:01:00", "aa:bb:cc:00:02:00", "aa:bb:cc:00:03:00"])

    def test_first_row_and_member_marks(self):
        first, second, _third = self.devices
        self.assertEqual(first["Hostname"], "CORE-STK")
        self.assertEqual(first["OOB IP"], "192.0.2.10")                    # Gi0/0 管理埠優先於 SVI
        self.assertEqual(first["CPU Usage"], "7%")                         # 無管線版本優先於「| i CPU」
        self.assertIn("(46.1%)", first["Memory Usage"])                    # platform sorted
        self.assertEqual((second["Hostname"], second["OOB IP"], second["CPU Usage"]), ("○ (Stack)", "○", "○"))
        self.assertEqual(second["System Log"], "")

    def test_environment(self):
        r = self.devices[0]
        self.assertEqual(r["Fan Status"], "正常")                          # NOT PRESENT 不算異常
        self.assertEqual(r["Power Status"], "異常")                        # 3B No Input Power / Bad

    def test_syslog_dedup_and_filter(self):
        log = self.devices[0]["System Log"].splitlines()
        self.assertEqual(len(log), 2)                                      # LINK-3-UPDOWN 排除、SYS-5 等級過濾
        self.assertIn("PS_FAIL", log[0])
        self.assertTrue(log[1].startswith("000104: Oct 22 09:30:00.456"))  # 重複訊息保留最後一次


class TestNexus(unittest.TestCase):
    """N9K 單機：show environment 格式依官方 NX-API CLI Reference 與 PSU Shutdown 範例。"""

    def setUp(self):
        self.devices, _, self.results = analyze("198.51.100.20@N9K-LEAF.log")
        self.r = self.devices[0]

    def test_basic_fields(self):
        self.assertEqual(self.results[0].dev_type, "NX-OS")
        self.assertEqual((self.r["Hostname"], self.r["Model"], self.r["Serial Number"]),
                         ("N9K-LEAF", "N9K-C93180YC-FX", "FDO00000TST"))
        self.assertEqual(self.r["Firmware Version"], "10.3(6)")
        self.assertEqual(self.r["OOB IP"], "198.51.100.20")                # mgmt0
        self.assertEqual(self.r["Ethernet MAC"], "aa:bb:cc:00:20:01")      # vPC local system-mac
        self.assertEqual(self.r["CPU Usage"], "17.5%")                     # 無 five minutes → user + kernel
        self.assertIn("(63.5%)", self.r["Memory Usage"])

    def test_psu_shutdown(self):
        self.assertEqual(self.r["Power Status"], "異常")                   # V22：Shutdown 判為異常
        self.assertEqual(self.r["Fan Status"], "正常")                     # Fan_in_PS2 歸類為電源

    def test_fan_failure_or_missing(self):
        log = ("N3K# show environment\nFan:\n---------------------------------------------------\n"
               "Fan             Model                Hw         Status\n"
               "---------------------------------------------------\n"
               "Fan-1           NXA-FAN-30CFM-F      --         ok\n"
               "Fan-2           NXA-FAN-30CFM-F      --         failure or missing\n"
               "N3K# exit\n")
        p = cla.NexusParser(log, source="n3k.log").run_all()
        self.assertEqual((p.fan_status, p.pwr_status), ("異常", "正常"))   # V21 的 \bFAIL\b 抓不到 failure

    def test_syslog_strips_vdc_tag(self):
        log = self.r["System Log"]
        self.assertNotIn("VDC-1", log)
        self.assertNotIn("DAEMON-3-SYSTEM_MSG", log)                       # 雜訊排除
        self.assertTrue(log.startswith("2026 Oct 21 10:00:01 N9K-LEAF %PLATFORM-2-PS_FAIL"))


class TestAireOS(unittest.TestCase):
    """5520：show sysinfo 欄位依 Cisco Wireless Controller Command Reference 8.10。"""

    def setUp(self):
        self.devices, self.aps, self.results = analyze("192.0.2.50@WLC5520.log")
        self.r = self.devices[0]

    def test_sysinfo_fields(self):
        self.assertEqual(self.results[0].dev_type, "AireOS")
        self.assertEqual((self.r["Hostname"], self.r["OOB IP"]), ("WLC5520", "192.0.2.50"))
        self.assertEqual((self.r["Model"], self.r["Serial Number"]), ("AIR-CT5520-K9", "FCH0000TST5"))
        self.assertEqual(self.r["Firmware Version"], "8.10.190.0")
        self.assertEqual(self.r["HA Role"], "Standalone")
        self.assertEqual(self.r["Ethernet MAC"], "AA:BB:CC:00:50:00")
        self.assertIn("(58.7%)", self.r["Memory Usage"])

    def test_cpu_falls_back_to_show_cpu(self):
        self.assertEqual(self.r["CPU Usage"], "3%")                        # sysinfo 無 CPU 欄 → show cpu

    def test_power_supply_lines(self):
        self.assertEqual(self.r["Fan Status"], "正常")
        self.assertEqual(self.r["Power Status"], "異常")                   # Power Supply 2: Present, Not OK

    def test_power_supply_absent_is_normal(self):
        text = read_fixture("192.0.2.50@WLC5520.log").replace("Present, Not OK", "Absent")
        p = cla.AireOSParser(text, source="x.log").run_all()
        self.assertEqual(p.pwr_status, "正常")

    def test_joined_aps_only(self):
        self.assertEqual([a["ap_name"] for a in self.aps], ["AP-LOBBY-01", "AP-OFFICE-02"])   # Not joined 排除
        lobby = self.aps[0]
        self.assertEqual((lobby["ap_model"], lobby["serial"]), ("AIR-AP2802I-H-K9", "FGL0000AP01"))
        self.assertEqual(lobby["eth_mac"], "aa:bb:cc:00:a1:01")
        self.assertEqual(lobby["firmware"], "8.10.190")                    # AP 版本沿用 WLC 版本前三段

    def test_syslog_newest_first(self):
        log = self.r["System Log"].splitlines()
        self.assertEqual(len(log), 2)
        self.assertIn("DISC_OTAP_ERR", log[0])
        self.assertIn("Oct 22 10:15:01.100", log[1])                       # 新 → 舊 反轉後保留最後一次


class TestC1300(unittest.TestCase):
    """C1300 兩台 Stack 的 show tech-support；show cpu utilization / show inventory 格式依官方 CLI Guide。"""

    def setUp(self):
        self.devices, _, self.results = analyze("C1300-ACCESS_show_tech.txt")

    def test_detect_and_stack(self):
        self.assertEqual(self.results[0].dev_type, "C1300")
        self.assertEqual([r["Serial Number"] for r in self.devices], ["PSZ0000TS01", "PSZ0000TS02"])
        self.assertEqual(self.devices[1]["Hostname"], "○ (Stack)")

    def test_fields(self):
        r = self.devices[0]
        self.assertEqual((r["Hostname"], r["OOB IP"], r["Firmware Version"]),
                         ("C1300-ACCESS", "192.0.2.30", "4.1.6.53"))            # Active-image，不是 Inactive
        self.assertEqual(r["Ethernet MAC"], "aa:bb:cc:00:13:00")
        self.assertEqual(r["CPU Usage"], "3%")
        self.assertEqual(r["Memory Usage"], "Total：1048576000 | Used：346030080 | Free：702545920 (33%)")
        self.assertEqual((r["Fan Status"], r["Power Status"]), ("異常", "正常"))    # Unit 2 風扇 Failure

    def test_syslog(self):
        log = self.devices[0]["System Log"].splitlines()
        self.assertEqual(len(log), 2)                                      # STP-W-PORTSTATUS 排除、I 等級不收
        self.assertIn("POE-E-OVERLOAD", log[0])
        self.assertTrue(log[1].startswith("22-Oct-2026 09:30:00 %AAA-W-REJECT"))


class TestArubaEnvironment(unittest.TestCase):
    """AOS-CX 電源 / 風扇狀態值依官方 show environment power-supply / fan 說明。"""

    def _parse(self, text):
        return cla.ArubaParser(text, source="aruba.txt").run_all()

    def test_psu_absent_slot_is_normal(self):
        text = read_fixture("aruba_vsf_show_tech.txt").replace(
            "1/1      N/A      N/A               OK                 AC     100V-240V  200",
            "1/1      N/A      N/A               OK                 AC     100V-240V  200\n"
            "1/2      N/A      N/A               Absent             --     --         0")
        self.assertEqual(self._parse(text).pwr_status, "正常")

    def test_psu_warning_and_fault(self):
        base = read_fixture("aruba_vsf_show_tech.txt")
        for state in ("Warning", "Input Fault", "Output Fault"):
            text = base.replace("N/A               OK      ", f"N/A               {state:<8}")
            self.assertEqual(self._parse(text).pwr_status, "異常", state)

    def test_fan_tray_fault(self):
        base = read_fixture("aruba_vsf_show_tech.txt").replace("left-to-back   fault   0", "left-to-back   ok      4000")
        self.assertEqual(self._parse(base).fan_status, "正常")
        tray = base.replace("show environment fan\n",
                            "show environment fan\n\nFan tray information\n"
                            "---------------------------------------------------------\n"
                            "Mbr/Tray  Description            Status  Serial Number  Fans\n"
                            "---------------------------------------------------------\n"
                            "1/1       JL000A Fan tray        fault   SG0000FT01     4\n")
        self.assertEqual(self._parse(tray).fan_status, "異常")


class TestExcel(unittest.TestCase):
    def _sheets(self, *names):
        devices, aps, _ = analyze(*names)
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "out.xlsx")
            cla.write_excel(path, devices, aps)
            wb = load_workbook(path)
            return {ws.title: [c.value for c in ws[1]] for ws in wb.worksheets}

    def test_vendor_sheets_created_only_when_detected(self):
        sheets = self._sheets("10.0.0.10@WLC-HA.log")
        self.assertEqual(list(sheets), ["Report", "AP Info", "WLC Info", "Switch Info"])
        sheets = self._sheets("aruba_vsf_show_tech.txt", "FGT-BRANCH.conf")
        self.assertIn("Aruba Info", sheets)
        self.assertIn("Fortigate Info", sheets)
        for name in ("AP Info", "WLC Info", "Switch Info", "Aruba Info", "Fortigate Info"):
            self.assertEqual(sheets[name][-1], "Source File")

    def test_all_fixtures_rows_by_sheet(self):
        names = sorted(n for n in os.listdir(FIXTURES) if n.lower().endswith(cla.LOG_EXTS))
        devices, aps, results = analyze(*names)
        self.assertFalse([r.source for r in results if r.error or r.dev_type == "Unknown"])
        sheets = cla.build_sheets(devices, aps)
        self.assertEqual(list(sheets), ["Report", "AP Info", "WLC Info", "Switch Info", "Aruba Info", "Fortigate Info"])
        count = lambda name: len(sheets[name][1])
        self.assertEqual(count("WLC Info"), 3)                             # C9800 HA ×2 + AireOS
        self.assertEqual(count("Switch Info"), 3 + 1 + 2 + 3)              # IOS Stack、NX-OS、C1300 Stack、Aruba VSF
        self.assertEqual(count("AP Info"), 3)
        self.assertEqual(count("Report"), len(devices) + len(aps))
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "all.xlsx")
            cla.write_excel(path, devices, aps)
            wb = load_workbook(path)
            longest = max(len(str(c.value)) for ws in wb.worksheets for row in ws.iter_rows() for c in row
                          if c.value is not None)
            self.assertLess(longest, 32767)


if __name__ == "__main__":
    unittest.main()
