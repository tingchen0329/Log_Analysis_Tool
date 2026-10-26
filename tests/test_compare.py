"""
維護設備比對測試（V23）：只使用 tests/compare_scenario.py 的合成資料與 tests/fixtures 的合成 Log。

執行：python -m unittest discover -s tests -v
"""
import contextlib
import importlib.util
import io
import os
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIXTURES = os.path.join(ROOT, "tests", "fixtures")
sys.path.insert(0, os.path.join(ROOT, "tests"))
import compare_scenario  # noqa: E402

_spec = importlib.util.spec_from_file_location("cla", os.path.join(ROOT, "Cisco_Log_Analysis.py"))
cla = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(cla)


def by(rows, **match):
    return [r for r in rows if all(r.get(k) == v for k, v in match.items())]


class TestCompareScenario(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.q2 = os.path.join(cls.tmp.name, "2026Q2.xlsx")
        cls.q3 = os.path.join(cls.tmp.name, "2026Q3.xlsx")
        cla.write_excel(cls.q2, *compare_scenario.quarter("Q2"))
        cla.write_excel(cls.q3, *compare_scenario.quarter("Q3"))
        cls.result = cla.compare_reports(cls.q2, cls.q3)
        cls.s = cls.result.sheets

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_counts(self):
        self.assertEqual(self.result.counts(), {
            "新增設備": 3, "移除設備": 1, "換機與成員變更": 2, "欄位異動": 6,   # 欄位異動以台數計
            "新增 AP": 2, "移除 AP": 1, "需人工確認": 1})
        self.assertEqual(len(self.s["欄位異動"]), 6)                     # 明細 6 列（成員 1、3 合併）
        self.assertEqual(self.result.device_stats(),
                         {"相符": 2, "換機": 1, "移除": 1, "新增": 3, "異動": 5})
        self.assertEqual(self.result.ap_stats(), {"新增": 2, "移除": 1, "換機": 1, "改名": 1})

    def test_mark_rows_filled(self):
        stack = [d for d in self.result.new.devices if d.host == "CORE-STK"]
        self.assertEqual([d.pos for d in stack], ["成員 1", "成員 2", "成員 3", "成員 4"])
        self.assertTrue(all(d.values["OOB IP"] == "192.0.2.10" for d in stack))
        ha = [d for d in self.result.new.devices if d.host == "WLC-HA"]
        self.assertEqual([d.pos for d in ha], ["機箱 1", "機箱 2"])

    def test_added_and_removed(self):
        self.assertEqual({r["Hostname"] for r in self.s["新增設備"]}, {"ACC-SW05", "FGT-BRANCH2", "CORE-STK"})
        self.assertEqual(by(self.s["新增設備"], Hostname="FGT-BRANCH2")[0]["類型"], "FortiGate")
        self.assertEqual([r["Serial Number"] for r in self.s["移除設備"]], ["FOC0000AS02"])

    def test_stack_member_replaced_and_added(self):
        rows = by(self.s["換機與成員變更"], Hostname="CORE-STK")
        replaced = by(rows, 判定="成員換新")[0]
        self.assertEqual((replaced["位置"], replaced["上季 Serial"], replaced["本季 Serial"]),
                         ("成員 2", "FOC0000SK02", "FOC0000SK05"))
        self.assertEqual(len(rows), 1)                             # Stack 新成員列在「新增設備」
        added = by(self.s["新增設備"], Hostname="CORE-STK")[0]
        self.assertEqual((added["位置"], added["Serial Number"]), ("成員 4", "FOC0000SK06"))

    def test_rename_is_not_add(self):
        self.assertFalse(by(self.s["新增設備"], Hostname="ACC-SW03-NEW"))
        row = by(self.s["欄位異動"], 欄位="Hostname（改名）")[0]
        self.assertEqual((row["上季"], row["本季"]), ("ACC-SW03", "ACC-SW03-NEW"))

    def test_field_changes(self):
        fw = by(self.s["欄位異動"], Hostname="CORE-STK", 欄位="Firmware Version")
        self.assertEqual(len(fw), 1)                               # 成員 1、3 相同異動合併為一列
        self.assertEqual(fw[0]["位置"], "成員 1、成員 3")
        pwr = by(self.s["欄位異動"], Hostname="N9K-LEAF", 欄位="Power Status")[0]
        self.assertEqual((pwr["上季"], pwr["本季"]), ("正常", "異常"))
        alert = by(self.s["欄位異動"], Hostname="N9K-LEAF", 欄位="System Log（新增告警）")[0]
        self.assertIn("PS_FAIL", alert["本季"])
        ip = by(self.s["欄位異動"], Hostname="FGT-HQ", 欄位="OOB IP")[0]
        self.assertEqual((ip["上季"], ip["本季"]), ("198.51.100.1", "198.51.100.2"))
        self.assertFalse(by(self.s["欄位異動"], Hostname="WLC-HA"))  # 沒變化的 HA 不列出

    def test_aps(self):
        self.assertEqual({a["AP Name"] for a in self.s["新增 AP"]}, {"AP-3F-07", "AP-3F-08"})
        self.assertEqual([a["AP Name"] for a in self.s["移除 AP"]], ["AP-OLD"])
        swap = by(self.s["換機與成員變更"], 判定="AP 換機")[0]
        self.assertEqual((swap["Hostname"], swap["上季 Serial"], swap["本季 Serial"]),
                         ("AP-2F-01", "FJC0000AP03", "FJC0000AP09"))
        self.assertTrue(by(self.s["欄位異動"], 欄位="AP Name（改名）", 本季="AP-1F-02B"))

    def test_no_serial_needs_review(self):
        row = self.s["需人工確認"][0]
        self.assertEqual((row["Hostname"], row["季別"]), ("EDGE-SW", "兩季"))

    def test_compare_excel(self):
        out = os.path.join(self.tmp.name, "cmp.xlsx")
        cla.write_compare_excel(out, self.result)
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, list(cla.COMPARE_SHEETS))
        self.assertEqual(wb.sheetnames[0], "設備總表")
        summary = {r[0]: r[1] for r in wb["摘要"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(summary["1. 序號完全相符（無變更）"], "2 台")
        self.assertEqual(summary["4. 新增設備 / Stack 新成員"], "3 台")
        self.assertEqual(summary["5. 欄位異動（以台數計）"], "5 台")
        self.assertIn("12 台設備", summary["本季報表"])
        self.assertEqual(wb["新增設備"].cell(1, 1).value, "類型")
        self.assertEqual(wb["新增設備"].cell(2, 1).fill.fgColor.rgb[-6:], cla.FILL_ADD)

    def test_master_sheet_colors(self):
        out = os.path.join(self.tmp.name, "master.xlsx")
        cla.write_compare_excel(out, self.result)
        ws = load_workbook(out)["設備總表"]
        header = [c.value for c in ws[3]]
        self.assertEqual(header, cla.MASTER_COLS)
        self.assertEqual((ws.freeze_panes, ws.auto_filter.ref), ("D4", f"A3:P{ws.max_row}"))
        col = {name: i + 1 for i, name in enumerate(header)}
        rows = {}
        for r in range(4, ws.max_row + 1):
            key = (ws.cell(r, col["Serial Number"]).value, ws.cell(r, col["狀態"]).value)
            rows[key] = r
        color = lambda r, name: ws.cell(r, col[name]).fill.fgColor.rgb[-6:]

        r = rows[("FOC0000SK01", "異動")]                         # Stack 首列：Firmware 黃 + 註解
        self.assertEqual(color(r, "Firmware Version"), cla.FILL_CHANGED)
        self.assertEqual(ws.cell(r, col["Firmware Version"]).comment.text, "上季：17.9.4")
        self.assertEqual(color(r, "Model"), "000000")
        r = rows[("FOC0000SK05", "換機")]                         # 成員換新：只標 SN 深紅
        self.assertEqual(ws.cell(r, col["Hostname"]).value, "○ (Stack)")
        self.assertEqual(color(r, "Serial Number"), cla.FILL_RED)
        self.assertEqual(color(r, "Model"), "000000")
        r = rows[("FOC0000SK06", "新增")]                         # Stack 新成員：整列淡綠
        self.assertEqual({color(r, n) for n in header}, {cla.FILL_ADD})
        r = rows[("N/A", "需確認")]
        self.assertEqual((color(r, "狀態"), color(r, "Serial Number")), (cla.FILL_REVIEW, cla.FILL_REVIEW))
        r = rows[("FDO0000NX01", "異動")]
        self.assertEqual(color(r, "Power Status"), cla.FILL_CHANGED)
        self.assertIn("PS_FAIL", ws.cell(r, col["System Log"]).comment.text)
        self.assertIsNone(ws.cell(rows[("FCL0000HA01", None)], col["Hostname"]).comment)   # 無變化：無標註

        removed = [r for (sn, st), r in rows.items() if st == "移除"]
        self.assertEqual(sorted(removed), [ws.max_row - 1, ws.max_row])   # 移除的設備在最下方
        for r in removed:
            self.assertEqual(color(r, "Model"), cla.FILL_REMOVED)
            self.assertEqual((color(r, "Hostname"), color(r, "Serial Number")), (cla.FILL_RED, cla.FILL_RED))
        self.assertEqual(ws.cell(ws.max_row - 1, col["Hostname"]).value, "ACC-SW02")

    def test_compare_items(self):
        items = cla.compare_items(self.result)
        self.assertEqual(len(items), 16)
        self.assertEqual(items[0][0], "新增")                     # 依新增 → 移除 → 換機 → 異動 → 需確認
        self.assertEqual(items[-1], ("需確認", "EDGE-SW", "無序號，以 Hostname + Model 配對"))
        self.assertIn(("換機", "CORE-STK 成員 2", "成員 2 換新：FOC0000SK02 → FOC0000SK05"), items)

    def test_reverse_direction(self):
        rev = cla.compare_reports(self.q3, self.q2).counts()
        self.assertEqual((rev["新增設備"], rev["移除設備"], rev["新增 AP"], rev["移除 AP"]), (1, 3, 1, 2))   # 移除含 Stack 成員 4

    def test_committed_example_files_match_scenario(self):
        q2, q3 = (os.path.join(FIXTURES, f"compare_2026{q}.xlsx") for q in ("Q2", "Q3"))
        self.assertEqual(cla.compare_reports(q2, q3).counts(), self.result.counts())


class TestCompareEdgeCases(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def path(self, name):
        return os.path.join(self.tmp.name, name)

    def test_identical_reports_from_real_parser(self):
        devices, aps, _ = cla.analyze_folder(FIXTURES, log=lambda _m: None)
        cla.write_excel(self.path("a.xlsx"), devices, aps)
        counts = cla.compare_reports(self.path("a.xlsx"), self.path("a.xlsx")).counts()
        self.assertEqual({k: v for k, v in counts.items() if k != "需人工確認"},
                         {k: 0 for k in counts if k != "需人工確認"})

    def test_stack_shared_fields_only_on_head(self):
        """Stack 改名 / 換管理 IP：成員列顯示「○」，只在首列標註，不重複列出。"""
        q2, aps = compare_scenario.quarter("Q2")
        q3 = [dict(d) for d in q2]
        q3[0].update({"Hostname": "CORE-STK-NEW", "OOB IP": "192.0.2.11"})
        cla.write_excel(self.path("a.xlsx"), q2, aps)
        cla.write_excel(self.path("b.xlsx"), q3, aps)
        r = cla.compare_reports(self.path("a.xlsx"), self.path("b.xlsx"))
        self.assertEqual([(c["欄位"], c["位置"]) for c in r.sheets["欄位異動"]],
                         [("Hostname（改名）", "成員 1"), ("OOB IP", "成員 1")])
        self.assertEqual(r.counts()["欄位異動"], 1)
        self.assertEqual([m.label for m in r.master if m.mark.status == "異動"], ["CORE-STK-NEW 成員 1"])

    def test_empty_sheets_written_as_none(self):
        devices, aps = compare_scenario.quarter("Q2")
        cla.write_excel(self.path("a.xlsx"), devices, aps)
        out = self.path("cmp.xlsx")
        cla.write_compare_excel(out, cla.compare_reports(self.path("a.xlsx"), self.path("a.xlsx")))
        self.assertEqual(load_workbook(out)["新增設備"].cell(2, 1).value, "無")

    def test_not_a_tool_report(self):
        wb = Workbook()
        wb.active.append(["Name", "Value"])
        wb.save(self.path("other.xlsx"))
        with self.assertRaises(cla.CompareError):
            cla.compare_reports(self.path("other.xlsx"), self.path("other.xlsx"))

    def test_not_an_excel_file(self):
        with open(self.path("bad.xlsx"), "w") as f:
            f.write("not a zip")
        with self.assertRaises(cla.CompareError):
            cla.load_report(self.path("bad.xlsx"))

    def test_duplicate_serial_flagged(self):
        devices, aps = compare_scenario.quarter("Q2")
        devices.append(dict(devices[-1], Hostname="FGT-DUP"))       # 同序號出現兩次
        cla.write_excel(self.path("a.xlsx"), devices, aps)
        review = cla.compare_reports(self.path("a.xlsx"), self.path("a.xlsx")).sheets["需人工確認"]
        self.assertTrue(any("序號重複" in r["原因"] for r in review))

    def test_alert_lines_ignore_timestamp(self):
        a = cla._alert_lines("Jun 10 08:00:01: %PLATFORM-2-PS_FAIL: Power supply 2 failed")
        b = cla._alert_lines("Sep 01 21:13:44: %PLATFORM-2-PS_FAIL: Power supply 2 failed")
        self.assertEqual(a.keys(), b.keys())
        self.assertEqual(cla._alert_lines("%LINK-3-UPDOWN: Interface Gi1/0/1, changed state to up"), {})


class TestCompareCli(unittest.TestCase):
    def test_compare_and_compare_with(self):
        with tempfile.TemporaryDirectory() as d:
            q2, q3 = os.path.join(d, "Q2.xlsx"), os.path.join(d, "Q3.xlsx")
            cla.write_excel(q2, *compare_scenario.quarter("Q2"))
            cla.write_excel(q3, *compare_scenario.quarter("Q3"))
            out = os.path.join(d, "cmp.xlsx")
            with contextlib.redirect_stdout(io.StringIO()) as buf:
                cla.main(["--compare", q2, q3, "-o", out])
                cla.main([FIXTURES, "-o", os.path.join(d, "now.xlsx"), "--compare-with", q2])
            self.assertIn("新增設備 3 台｜移除設備 1 台", buf.getvalue())
            self.assertIn("欄位異動", load_workbook(out).sheetnames)

            report = os.path.join(d, "now.xlsx")
            self.assertTrue(os.path.exists(report))
            self.assertTrue(os.path.exists(os.path.join(d, cla.default_compare_name(q2, report))))


if __name__ == "__main__":
    unittest.main()
